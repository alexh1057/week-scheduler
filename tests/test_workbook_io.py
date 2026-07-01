import math

import pytest
from openpyxl import load_workbook

from structural2d.excel.build_template import build_template
from structural2d.excel.workbook_io import read_model_openpyxl, write_results_openpyxl
from structural2d.sections import section_lookup
from structural2d.solver import AnalysisError, solve


def test_template_round_trip(tmp_path):
    """build_template's own sample data (simply supported beam, midspan
    point load) should read back, solve, and write results identically to
    the closed-form solution, and the Diagrams sheet should get images."""
    path = tmp_path / "frame_model.xlsx"
    build_template(str(path))

    wb = load_workbook(str(path))
    model = read_model_openpyxl(wb)

    assert set(model.nodes) == {"1", "2"}
    assert model.members["m1"].node_i == "1"
    assert model.members["m1"].node_j == "2"
    assert model.supports["1"].ux and model.supports["1"].uy and not model.supports["1"].rz
    assert not model.supports["2"].ux and model.supports["2"].uy
    assert len(model.point_loads) == 1

    results = solve(model)
    P, L = 10_000.0, 6.0
    rx1, ry1, _ = results.reactions["1"]
    rx2, ry2, _ = results.reactions["2"]
    assert math.isclose(ry1, P / 2, rel_tol=1e-9)
    assert math.isclose(ry2, P / 2, rel_tol=1e-9)

    write_results_openpyxl(wb, model, results)
    wb.save(str(path))

    wb2 = load_workbook(str(path))
    reaction_rows = list(wb2["Reactions"].iter_rows(values_only=True))
    assert reaction_rows[0] == ("node_id", "Rx", "Ry", "Rm")
    assert reaction_rows[1][0] == "1"
    assert math.isclose(reaction_rows[1][2], P / 2, rel_tol=1e-9)

    displacement_rows = list(wb2["Displacements"].iter_rows(values_only=True))
    assert displacement_rows[0] == ("node_id", "Ux", "Uy", "Rz")

    assert len(wb2["Diagrams"]._images) == 3


def test_rerunning_write_results_does_not_duplicate_images(tmp_path):
    """write_results_openpyxl clears the Diagrams sheet first, so calling
    it twice (as happens every time the user re-runs the analysis) must
    not accumulate images or stale table rows."""
    path = tmp_path / "frame_model.xlsx"
    build_template(str(path))
    wb = load_workbook(str(path))
    model = read_model_openpyxl(wb)
    results = solve(model)

    write_results_openpyxl(wb, model, results)
    write_results_openpyxl(wb, model, results)

    assert len(wb["Diagrams"]._images) == 3
    reaction_rows = [r for r in wb["Reactions"].iter_rows(values_only=True) if r[0] is not None]
    assert len(reaction_rows) == 3  # header + 2 nodes, no leftover duplicates


def test_section_pick_fills_blank_EAI(tmp_path):
    """Picking a section from the dropdown and leaving E/A/I blank should
    resolve the member properties from the Sections sheet."""
    path = tmp_path / "frame_model.xlsx"
    build_template(str(path))
    wb = load_workbook(str(path))
    ws = wb["Members"]
    ws["D2"] = "305x165x40 UB"
    ws["E2"] = ws["F2"] = ws["G2"] = None

    model = read_model_openpyxl(wb)
    E, A, I = section_lookup()["305x165x40 UB"]
    m = model.members["m1"]
    assert (m.E, m.A, m.I) == (E, A, I)
    solve(model)  # and the model still solves


def test_unknown_section_raises_readable_error(tmp_path):
    path = tmp_path / "frame_model.xlsx"
    build_template(str(path))
    wb = load_workbook(str(path))
    ws = wb["Members"]
    ws["D2"] = "999x999x9 UB"
    ws["E2"] = ws["F2"] = ws["G2"] = None
    with pytest.raises(ValueError, match="not on the .*Sections"):
        read_model_openpyxl(wb)


def test_summary_sheet_written(tmp_path):
    path = tmp_path / "frame_model.xlsx"
    build_template(str(path))
    wb = load_workbook(str(path))
    model = read_model_openpyxl(wb)
    results = solve(model)
    write_results_openpyxl(wb, model, results)

    rows = [r for r in wb["Summary"].iter_rows(values_only=True) if r[0] is not None]
    assert rows[0] == ("member_id", "max_abs_N", "max_abs_V", "max_abs_M", "max_abs_deflection")
    assert rows[1][0] == "m1"
    P, L = 10_000.0, 6.0
    assert math.isclose(rows[1][2], P / 2, rel_tol=1e-6)  # max |V|
    assert math.isclose(rows[1][3], P * L / 4, rel_tol=1e-6)  # max |M|


def test_dangling_node_reference_raises_readable_error(tmp_path):
    path = tmp_path / "frame_model.xlsx"
    build_template(str(path))
    wb = load_workbook(str(path))
    wb["Members"]["C2"] = "99"  # node_j that doesn't exist
    model = read_model_openpyxl(wb)
    with pytest.raises(AnalysisError, match="refers to node '99'"):
        solve(model)
