"""Headless validation of the two-cell split engine:
  engine_solver.py  -> Engine!B2 (no matplotlib)
  engine_plotter.py -> Engine!B4 (matplotlib only, takes solver output)
pandas DataFrames stand in for xl(...) results.
"""
import math
import importlib.util, pathlib

import pandas as pd
import pytest

from structural2d.excel.build_pyexcel_template import _plotter_paste_text, _solver_paste_text

# Load the modules from their source files (they're not importable as a
# package since they're meant to be pasted verbatim into Excel cells)
def _load(name):
    path = pathlib.Path(__file__).parent.parent / "structural2d" / "pyexcel" / f"{name}.py"
    spec = importlib.util.spec_from_file_location(name, path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod

solver = _load("engine_solver")
plotter = _load("engine_plotter")

PYEXCEL_CELL_LIMIT = 8192  # Microsoft's hard per-cell limit for Insert Python


def _ssb_tables():
    """Simply supported beam, midspan 10 kN point load."""
    nodes = pd.DataFrame([["1", 0.0, 0.0], ["2", 6.0, 0.0]], columns=["id", "x", "y"])
    members = pd.DataFrame(
        [["m1", "1", "2", 200e9, 0.01, 8e-5, "FALSE", "FALSE"]],
        columns=["id", "node_i", "node_j", "E", "A", "I", "hinge_i", "hinge_j"],
    )
    supports = pd.DataFrame([["1", "Pinned"], ["2", "Roller Y"]], columns=["node_id", "type"])
    nodal_loads = pd.DataFrame(columns=["node_id", "fx", "fy", "m"])
    point_loads = pd.DataFrame(
        [["m1", 3.0, 0.0, -10000.0, 0.0, "local"]],
        columns=["member_id", "position", "fx", "fy", "m", "frame"],
    )
    udls = pd.DataFrame(columns=["member_id", "wx", "wy", "frame", "start", "end"])
    return nodes, members, supports, nodal_loads, point_loads, udls


def test_solver_reactions_match_closed_form():
    result = solver.run(*_ssb_tables())
    P = 10_000.0
    rx = result["reactions"].set_index("node_id")
    assert math.isclose(rx.loc["1", "Ry"], P / 2, rel_tol=1e-9)
    assert math.isclose(rx.loc["2", "Ry"], P / 2, rel_tol=1e-9)
    assert math.isclose(rx.loc["1", "Rx"], 0.0, abs_tol=1e-6)


def test_solver_end_rotation_matches_closed_form():
    result = solver.run(*_ssb_tables())
    P, L, E, I = 10_000.0, 6.0, 200e9, 8e-5
    theta = -P * L**2 / (16 * E * I)
    dx = result["displacements"].set_index("node_id")
    assert math.isclose(dx.loc["1", "Rz"], theta, rel_tol=1e-4)


def test_solver_output_keys():
    result = solver.run(*_ssb_tables())
    for key in ("reactions", "displacements", "nodes", "members", "supports",
                "nlx", "plx", "ulx", "samples", "deformed", "deformed_scale"):
        assert key in result, f"missing key: {key}"


def test_solver_blank_optional_columns():
    nodes, members, supports, nodal_loads, point_loads, udls = _ssb_tables()
    members = pd.DataFrame(
        [["m1", "1", "2", 200e9, 0.01, 8e-5, None, None]],
        columns=["id", "node_i", "node_j", "E", "A", "I", "hinge_i", "hinge_j"],
    )
    point_loads = pd.DataFrame(
        [["m1", 3.0, 0.0, -10000.0, 0.0, None]],
        columns=["member_id", "position", "fx", "fy", "m", "frame"],
    )
    result = solver.run(nodes, members, supports, nodal_loads, point_loads, udls)
    rx = result["reactions"].set_index("node_id")
    assert math.isclose(rx.loc["1", "Ry"], 5000.0, rel_tol=1e-9)


def test_solver_blank_placeholder_rows_skipped():
    nodes, members, supports, _, point_loads, _ = _ssb_tables()
    nodal_loads = pd.DataFrame([[None, None, None, None]], columns=["node_id", "fx", "fy", "m"])
    udls = pd.DataFrame([[None, None, None, None, None, None]],
                        columns=["member_id", "wx", "wy", "frame", "start", "end"])
    result = solver.run(nodes, members, supports, nodal_loads, point_loads, udls)
    rx = result["reactions"].set_index("node_id")
    assert math.isclose(rx.loc["1", "Ry"], 5000.0, rel_tol=1e-9)


def test_plotter_returns_figures():
    d = solver.run(*_ssb_tables())
    out = plotter.run(d)
    assert out["geometry_fig"] is not None
    assert out["deformed_fig"] is not None
    fig = out["member_fig"]("m1")
    assert fig is not None


def test_plotter_with_nodal_and_udl_loads():
    """Plotter must render without error when all load types are present."""
    nodes = pd.DataFrame([["1", 0.0, 0.0], ["2", 6.0, 0.0]], columns=["id", "x", "y"])
    members = pd.DataFrame(
        [["m1", "1", "2", 200e9, 0.01, 8e-5, "FALSE", "FALSE"]],
        columns=["id", "node_i", "node_j", "E", "A", "I", "hinge_i", "hinge_j"],
    )
    supports = pd.DataFrame([["1", "Fixed"]], columns=["node_id", "type"])
    nodal_loads = pd.DataFrame([["2", 5000.0, -2000.0, 0.0]], columns=["node_id", "fx", "fy", "m"])
    point_loads = pd.DataFrame(columns=["member_id", "position", "fx", "fy", "m", "frame"])
    udls = pd.DataFrame([["m1", 0.0, -1000.0, "global", None, None]],
                        columns=["member_id", "wx", "wy", "frame", "start", "end"])
    d = solver.run(nodes, members, supports, nodal_loads, point_loads, udls)
    out = plotter.run(d)
    assert out["geometry_fig"] is not None


MEMBER_COLS = ["id", "node_i", "node_j", "section", "E", "A", "I", "hinge_i", "hinge_j"]
SECTION_ROWS = pd.DataFrame(
    [["305x165x40 UB", 210e9, 51.3e-4, 8503e-8]], columns=["name", "E", "A", "I"]
)


def _ssb_with_section(section, E=None, A=None, I=None):
    """Same beam, but with the new section-column Members layout."""
    nodes, _, supports, nodal_loads, point_loads, udls = _ssb_tables()
    members = pd.DataFrame(
        [["m1", "1", "2", section, E, A, I, "FALSE", "FALSE"]], columns=MEMBER_COLS
    )
    return nodes, members, supports, nodal_loads, point_loads, udls, SECTION_ROWS


def test_section_pick_fills_blank_EAI():
    result = solver.run(*_ssb_with_section("305x165x40 UB"))
    m = result["members"]["m1"]
    assert m["E"] == 210e9 and m["A"] == 51.3e-4 and m["I"] == 8503e-8
    # and the model actually solves with those properties
    rx = result["reactions"].set_index("node_id")
    assert math.isclose(rx.loc["1", "Ry"], 5000.0, rel_tol=1e-9)


def test_typed_EAI_overrides_section_pick():
    result = solver.run(*_ssb_with_section("305x165x40 UB", E=200e9, A=0.01, I=8e-5))
    m = result["members"]["m1"]
    assert m["E"] == 200e9 and m["A"] == 0.01 and m["I"] == 8e-5


def test_unknown_section_raises_friendly_error():
    with pytest.raises(RuntimeError, match="unknown section"):
        solver.run(*_ssb_with_section("999x999x9 UB"))


def test_blank_section_and_blank_EAI_raises_friendly_error():
    with pytest.raises(RuntimeError, match="pick a section or fill in E/A/I"):
        solver.run(*_ssb_with_section(None))


def test_unknown_node_reference_raises_friendly_error():
    nodes, members, supports, nodal_loads, point_loads, udls = _ssb_tables()
    members = pd.DataFrame(
        [["m1", "1", "99", 200e9, 0.01, 8e-5, "FALSE", "FALSE"]],
        columns=["id", "node_i", "node_j", "E", "A", "I", "hinge_i", "hinge_j"],
    )
    with pytest.raises(RuntimeError, match="unknown node"):
        solver.run(nodes, members, supports, nodal_loads, point_loads, udls)


def test_unknown_member_on_load_raises_friendly_error():
    nodes, members, supports, nodal_loads, point_loads, udls = _ssb_tables()
    point_loads = pd.DataFrame(
        [["m99", 3.0, 0.0, -10000.0, 0.0, "local"]],
        columns=["member_id", "position", "fx", "fy", "m", "frame"],
    )
    with pytest.raises(RuntimeError, match="unknown member"):
        solver.run(nodes, members, supports, nodal_loads, point_loads, udls)


def test_summary_matches_closed_form():
    result = solver.run(*_ssb_tables())
    s = result["summary"].set_index("member_id")
    P, L, E, I = 10_000.0, 6.0, 200e9, 8e-5
    assert math.isclose(s.loc["m1", "max_abs_V"], P / 2, rel_tol=1e-6)
    assert math.isclose(s.loc["m1", "max_abs_M"], P * L / 4, rel_tol=1e-6)
    # midspan deflection P*L^3 / (48*E*I)
    assert math.isclose(s.loc["m1", "max_abs_deflection"], P * L**3 / (48 * E * I), rel_tol=1e-3)


def test_staged_paste_text_under_excel_cell_limit():
    """_solver_paste_text()/_plotter_paste_text() are exactly what
    build_pyexcel_template stages in Engine!D2/D4 for the user to
    copy-paste into the Insert Python cells B2/B4 -- each must stay under
    Excel's hard 8192-character per-cell limit, source + run() line
    included, or the one-click paste breaks."""
    assert len(_solver_paste_text()) < PYEXCEL_CELL_LIMIT
    assert len(_plotter_paste_text()) < PYEXCEL_CELL_LIMIT


def test_staged_solver_paste_text_executes_and_matches_closed_form():
    """The exact string staged in Engine!D2 -- source plus the appended
    run() line -- must itself be valid, runnable Python that reproduces
    the same closed-form result, since that's what the user pastes
    verbatim (not the .py file's contents plus manual typing)."""
    code = _solver_paste_text()
    body = "\n".join(code.splitlines()[:-1])  # drop the trailing run() line
    ns = {}
    exec(body, ns)
    nodes, members, supports, nodal_loads, point_loads, udls = _ssb_tables()
    sections = pd.DataFrame(columns=["name", "E", "A", "I"])
    result = ns["run"](nodes, members, supports, nodal_loads, point_loads, udls, sections)
    rx = result["reactions"].set_index("node_id")
    assert math.isclose(rx.loc["1", "Ry"], 5000.0, rel_tol=1e-9)


def test_staged_plotter_paste_text_executes_and_returns_figures():
    code = _plotter_paste_text()
    body = "\n".join(code.splitlines()[:-1])
    ns = {}
    exec(body, ns)
    d = solver.run(*_ssb_tables())
    out = ns["run"](d)
    assert out["geometry_fig"] is not None
    assert out["member_fig"]("m1") is not None


def test_workbook_stages_the_same_paste_text_in_engine_D2_D4(tmp_path):
    """build_pyexcel_template must actually write the staged text into the
    workbook's Engine!D2/D4 cells (not just make it available as a helper
    function) -- that's the whole point of the staging-cell workflow."""
    from openpyxl import load_workbook

    from structural2d.excel.build_pyexcel_template import build_pyexcel_template

    path = tmp_path / "frame_model_pyexcel.xlsx"
    build_pyexcel_template(str(path))
    wb = load_workbook(str(path))
    ws = wb["Engine"]
    assert ws["D2"].value == _solver_paste_text()
    assert ws["D4"].value == _plotter_paste_text()
    assert ws["B2"].value is None  # left empty for the user's Insert Python
    assert ws["B4"].value is None
