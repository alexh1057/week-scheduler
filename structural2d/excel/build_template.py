"""Scaffolds the input/output workbook for the 2D frame analysis tool.

Pure openpyxl (no Excel/xlwings needed), so it's testable headlessly. The
generated workbook is later opened in Excel and read/written live by
workbook_io.py through the xlwings "Run main" button (see frame.py).
"""
from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from ..sections import SECTION_HEADERS, SECTIONS

_HEADER_FILL = PatternFill(start_color="FFD9E1F2", end_color="FFD9E1F2", fill_type="solid")
_HEADER_FONT = Font(bold=True)

SUPPORT_TYPES = ["Fixed", "Pinned", "Roller X", "Roller Y"]
LOAD_FRAMES = ["local", "global"]
BOOL_CHOICES = ["TRUE", "FALSE"]

# Rows the section-name dropdown scans on the Sections sheet: covers the
# built-in catalog plus plenty of headroom for user-added sections.
SECTION_DROPDOWN_ROWS = 100


def _add_section_name(wb: Workbook) -> None:
    """Workbook-level defined name for the dropdown source -- a named range
    works in every Excel version, unlike direct cross-sheet DV references."""
    wb.defined_names.add(
        DefinedName("SectionNames", attr_text=f"Sections!$A$2:$A${SECTION_DROPDOWN_ROWS + 1}")
    )


def _add_section_dropdown(ws: Worksheet, col_letter: str, n_rows: int = 200) -> None:
    dv = DataValidation(type="list", formula1="=SectionNames", allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{col_letter}2:{col_letter}{n_rows + 1}")


def _write_header(ws: Worksheet, headers: list[str]) -> None:
    for col, name in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=name)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"


def _autosize(ws: Worksheet, widths: list[int]) -> None:
    for col, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=col).column_letter].width = width


def _add_list_validation(ws: Worksheet, col_letter: str, choices: list[str], n_rows: int = 200) -> None:
    dv = DataValidation(type="list", formula1=f'"{",".join(choices)}"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"{col_letter}2:{col_letter}{n_rows + 1}")


def _build_nodes(wb: Workbook) -> None:
    ws = wb.create_sheet("Nodes")
    _write_header(ws, ["id", "x", "y"])
    ws.append(["1", 0.0, 0.0])
    ws.append(["2", 6.0, 0.0])
    _autosize(ws, [10, 12, 12])


def _build_members(wb: Workbook) -> None:
    ws = wb.create_sheet("Members")
    _write_header(ws, ["id", "node_i", "node_j", "section", "E", "A", "I", "hinge_i", "hinge_j"])
    ws.append(["m1", "1", "2", None, 200e9, 0.01, 8e-5, "FALSE", "FALSE"])
    _autosize(ws, [10, 10, 10, 18, 14, 12, 12, 10, 10])
    _add_section_dropdown(ws, "D")
    _add_list_validation(ws, "H", BOOL_CHOICES)
    _add_list_validation(ws, "I", BOOL_CHOICES)


def _build_sections(wb: Workbook) -> None:
    ws = wb.create_sheet("Sections")
    _write_header(ws, SECTION_HEADERS)
    for row in SECTIONS:
        ws.append(list(row))
    _autosize(ws, [18, 14, 12, 12])
    ws["F1"] = (
        "UK UB/UC sections, major-axis I, steel E=210e9 (SI units: Pa, m2, m4). "
        "Indicative values -- verify against current section tables for design "
        "use. Add your own rows; the Members dropdown picks them up."
    )
    ws.column_dimensions["F"].width = 110


def _build_supports(wb: Workbook) -> None:
    ws = wb.create_sheet("Supports")
    _write_header(ws, ["node_id", "type"])
    ws.append(["1", "Pinned"])
    ws.append(["2", "Roller Y"])
    _autosize(ws, [10, 14])
    _add_list_validation(ws, "B", SUPPORT_TYPES)


def _build_nodal_loads(wb: Workbook) -> None:
    ws = wb.create_sheet("NodalLoads")
    _write_header(ws, ["node_id", "fx", "fy", "m"])
    _autosize(ws, [10, 12, 12, 12])


def _build_point_loads(wb: Workbook) -> None:
    ws = wb.create_sheet("PointLoads")
    _write_header(ws, ["member_id", "position", "fx", "fy", "m", "frame"])
    ws.append(["m1", 3.0, 0.0, -10000.0, 0.0, "local"])
    _autosize(ws, [10, 12, 12, 12, 12, 10])
    _add_list_validation(ws, "F", LOAD_FRAMES)


def _build_udls(wb: Workbook) -> None:
    ws = wb.create_sheet("UDLs")
    _write_header(ws, ["member_id", "wx", "wy", "frame", "start", "end"])
    _autosize(ws, [10, 12, 12, 10, 12, 12])
    _add_list_validation(ws, "D", LOAD_FRAMES)


def _build_reactions(wb: Workbook) -> None:
    ws = wb.create_sheet("Reactions")
    _write_header(ws, ["node_id", "Rx", "Ry", "Rm"])
    _autosize(ws, [10, 14, 14, 14])


def _build_displacements(wb: Workbook) -> None:
    ws = wb.create_sheet("Displacements")
    _write_header(ws, ["node_id", "Ux", "Uy", "Rz"])
    _autosize(ws, [10, 14, 14, 14])


def _build_summary(wb: Workbook) -> None:
    ws = wb.create_sheet("Summary")
    _write_header(ws, ["member_id", "max_abs_N", "max_abs_V", "max_abs_M", "max_abs_deflection"])
    _autosize(ws, [12, 16, 16, 16, 20])


def _build_diagrams(wb: Workbook) -> None:
    wb.create_sheet("Diagrams")


def _build_instructions(wb: Workbook) -> None:
    ws = wb.create_sheet("Instructions", 0)
    ws.column_dimensions["A"].width = 100
    lines = [
        "2D Frame Analysis -- Instructions",
        "",
        "Units are not enforced -- use one consistent set throughout "
        "(e.g. N, m, Pa) across every sheet.",
        "",
        "Input sheets",
        "  Nodes        -- id, x, y",
        "  Members      -- id, node_i, node_j, section, E, A, I, hinge_i/hinge_j. "
        "EITHER pick a section from the dropdown and leave E/A/I blank (properties "
        "come from the Sections sheet), OR leave section blank and type E "
        "(modulus), A (area), I (moment of inertia) yourself. hinge_i/hinge_j "
        "TRUE releases that end's moment connection, e.g. for truss bars.",
        "  Sections     -- name, E, A, I: the catalog behind the Members "
        "dropdown. Ships with common UK UB/UC steel sections (SI units, "
        "indicative values -- verify for design use). Add rows for your own.",
        "  Supports     -- node_id, type (Fixed / Pinned / Roller X / Roller Y)",
        "  NodalLoads   -- node_id, fx, fy, m (force/moment applied directly at a node)",
        "  PointLoads   -- member_id, position (distance from node_i), fx, fy, m, "
        "frame (local: along the member axis: global: along the X/Y axes)",
        "  UDLs         -- member_id, wx, wy (force per length), frame, "
        "start/end (distance from node_i; leave blank for the full member)",
        "",
        "Output sheets (filled in automatically when you run the analysis)",
        "  Reactions, Displacements -- numeric results per node",
        "  Summary                  -- per-member max |N|, |V|, |M|, |deflection|",
        "  Diagrams                 -- geometry/load diagram, deformed shape, "
        "and per-member N/V/M diagrams",
        "",
        "Running the analysis",
        "  This workbook is paired with frame.py via xlwings. With the "
        "xlwings Excel add-in installed, click the xlwings ribbon tab's "
        "'Run main' button (or run `python frame.py` after editing it to "
        "call main() against this workbook) to solve the model and refresh "
        "the output sheets in place.",
        "",
        "The sample data already in the input sheets is a simply supported "
        "beam with a midspan point load -- replace it with your own model.",
    ]
    for row, text in enumerate(lines, start=1):
        ws.cell(row=row, column=1, value=text)
    ws.cell(row=1, column=1).font = Font(bold=True, size=14)
    for row, text in enumerate(lines, start=1):
        if text in ("Input sheets",
                    "Output sheets (filled in automatically when you run the analysis)",
                    "Running the analysis"):
            ws.cell(row=row, column=1).font = Font(bold=True)


def build_template(path: str) -> None:
    wb = Workbook()
    wb.remove(wb.active)
    _build_instructions(wb)
    _build_nodes(wb)
    _build_members(wb)
    _build_sections(wb)
    _build_supports(wb)
    _build_nodal_loads(wb)
    _build_point_loads(wb)
    _build_udls(wb)
    _build_reactions(wb)
    _build_displacements(wb)
    _build_summary(wb)
    _build_diagrams(wb)
    _add_section_name(wb)
    wb.save(path)


if __name__ == "__main__":
    import sys

    out_path = sys.argv[1] if len(sys.argv) > 1 else "frame_model.xlsx"
    build_template(out_path)
    print(f"Wrote {out_path}")
